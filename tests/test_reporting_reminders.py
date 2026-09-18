"""Filtered exports are real workbooks; reminders target assignees once per day."""
from datetime import date
from io import BytesIO
import tempfile
import unittest

from openpyxl import load_workbook
from test_server import app
import config
from database import insert_record
from relational_values import save_value
from reminders import deliver_due_reminders
from reports import report, report_csv, report_xls


class ReportingReminderTests(unittest.TestCase):
    def setUp(self):
        self.previous = config.DATA_DIR
        self.storage = tempfile.TemporaryDirectory(prefix="audit-reporting-")
        app.configure_data_directory(self.storage.name)
        app.init_db()

    def tearDown(self):
        app.configure_data_directory(self.previous)
        self.storage.cleanup()

    def test_filters_and_real_excel_and_room_trend(self):
        with app.connect() as db:
            for unit, outlet, day in (("Mini Studio", "A", "2026-01-02"), ("Mini Studio", "B", "2026-02-02"), ("Mini Studio", "A", "2026-02-03"), ("Loudspeaker", "A", "2026-01-02")):
                audit_id = insert_record(db, "audits", {"business_unit": unit, "outlet": outlet, "branch": "Room", "audit_date": day, "auditor": "Tester", "audit_type": "Standard", "score": 50, "created_at": 0})
                items = [{"location": "Room 1", "passed": True}, {"location": "Room 1", "passed": False}]
                insert_record(db, "inspection_sessions", {"business_unit": unit, "outlet": outlet, "zone": "Room", "audit_date": day, "auditor": "Tester", "items_data_id": save_value(db, items), "progress": 100, "status": "Completed", "audit_id": audit_id, "created_at": 0, "updated_at": 0})
                insert_record(db, "findings", {"finding_ref": f"F-{unit}-{outlet}-{day}", "audit_id": audit_id, "business_unit": unit, "outlet": outlet, "location": "Room", "category": "Safety", "priority": "Priority", "comment": "=1+1", "status": "Assigned", "created_at": 0, "updated_at": 0})
        filters = {"from": "2026-01-01", "to": "2026-01-31", "outlet": "A"}
        result = report("Mini Studio", filters)
        self.assertEqual(result["monthlySummary"]["auditsCompleted"], 1)
        self.assertEqual(result["monthlySummary"]["totalFindings"], 1)
        self.assertEqual(result["rankings"][0]["audit_date"], "2026-01-02")
        self.assertEqual(result["charts"]["roomAuditTrend"], [{"label": "A / Room 1 / 2026-01", "score": 50}])
        csv = report_csv("Mini Studio", filters).decode()
        self.assertIn("F-Mini Studio-A", csv)
        self.assertNotIn("F-Mini Studio-B", csv)
        self.assertNotIn("F-Loudspeaker-A", csv)
        self.assertIn("'=1+1", csv)
        workbook = load_workbook(BytesIO(report_xls("Mini Studio", filters)))
        self.assertEqual(workbook.sheetnames, ["Summary", "Findings"])
        self.assertEqual(workbook["Findings"].max_row, 2)
        self.assertEqual(workbook["Findings"]["M2"].value, "=1+1")
        self.assertEqual(workbook["Findings"]["M2"].data_type, "s")
        for invalid in ({"from": "wrong"}, {"from": "2026-02-01", "to": "2026-01-01"}):
            with self.assertRaises(ValueError):
                report("Mini Studio", invalid)

    def test_reminders_are_addressed_deduplicated_and_stop_on_completion(self):
        with app.connect() as db:
            user_id = insert_record(db, "users", {"name": "Assigned Tester", "email": "assigned@example.com", "role": "Department/PIC", "department": "TECH", "active": 1, "created_at": 0})
            common = {"business_unit": "Ottotree", "outlet": "STP", "zone": "Room", "request_type": "TECH", "priority": "Priority", "title": "Repair", "assignee": "Assigned Tester", "pic": "Assigned Tester", "status": "Assigned", "created_at": 0}
            due = insert_record(db, "work_orders", common | {"due_date": "2026-09-20"})
            overdue = insert_record(db, "work_orders", common | {"due_date": "2026-09-17"})
            insert_record(db, "work_orders", common | {"due_date": "2026-09-17", "status": "Completed"})
        self.assertEqual(deliver_due_reminders(date(2026, 9, 18)), 2)
        self.assertEqual(deliver_due_reminders(date(2026, 9, 18)), 0)
        with app.connect() as db:
            notices = db.execute("SELECT title, recipient_user_id FROM notifications WHERE related_type = 'work_order'").fetchall()
            self.assertEqual({row["title"] for row in notices}, {"Due soon", "Overdue"})
            self.assertEqual({row["recipient_user_id"] for row in notices}, {user_id})
            db.execute("UPDATE work_orders SET status = 'Completed' WHERE id IN (?,?)", (due, overdue))
        self.assertEqual(deliver_due_reminders(date(2026, 9, 19)), 0)
