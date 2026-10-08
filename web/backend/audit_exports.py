"""Audit exports beyond the summary report: one audit as Excel (an overall sheet, then a sheet
per location in full), one audit's chosen locations as PDF, and every audit of one location as
PDF or Excel."""
from io import BytesIO

from backend import config
from backend.accounts import branding_settings
from backend.database import connect
from backend.inspections import inspection_session
from backend.media_store import MediaStore
from backend.relational_values import load_value
from backend.scoring import summarize as summarize_score


def app_settings():
    with connect() as db:
        return {row["key"]: load_value(row["value_data_id"]) for row in db.execute("SELECT key, value_data_id FROM app_settings")}


def finish_workbook(workbook):
    """Bold headings, filters, readable widths; text that looks like a formula stays text."""
    from openpyxl.styles import Font
    for sheet in workbook:
        if sheet.max_row > 1:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            for cell in column:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def item_location(session, item):
    return item.get("location") or session.get("zone") or "Unassigned"


def session_locations(session):
    """The audit's locations, in the order they were inspected."""
    return list(dict.fromkeys(item_location(session, item) for item in session.get("items", [])))


def asset_records(items):
    """The register details of each asset checked (code, type, brand, model, serial, dates,
    status, photos), by its id as text: items carry the id as the page sent it."""
    ids = sorted({int(item["equipmentId"]) for item in items if str(item.get("equipmentId") or "").isdigit()})
    records = {}
    with connect() as db:
        for start in range(0, len(ids), 500):
            chunk = ids[start:start + 500]
            for row in db.execute(
                    "SELECT id, name, asset_id, code, kind, category, type, equipment_type, brand, model, serial_number, installation_date, "
                    "warranty_date, expiry_date, operational_status, health_status, location, zone, photos_data_id "
                    f"FROM equipment WHERE id IN ({','.join('?' for _ in chunk)})", chunk):
                record = dict(row)
                reference = record.pop("photos_data_id")
                record["photos"] = (load_value(reference) or []) if reference is not None else []
                record["type"] = record["type"] or record["equipment_type"] or ""
                record["status"] = record["operational_status"] or record["health_status"] or ""
                records[str(row["id"])] = record
    return records


def location_part(session, location, settings):
    """One location of an audit: its checks, findings, and score."""
    items = [item for item in session.get("items", []) if item_location(session, item) == location]
    findings = [finding for finding in session.get("findings") or [] if (finding.get("location") or session.get("zone")) == location]
    return {"location": location, "items": items, "findings": findings, "summary": summarize_score(items, settings),
            "assets": asset_records(items)}


def sheet_title(name, taken):
    """A worksheet name Excel accepts (31 characters, none of []:*?/\\), unique in the workbook."""
    base = "".join("-" if char in '[]:*?/\\' else char for char in str(name)).strip("' ")[:31] or "Location"
    title, number = base, 2
    while title.lower() in taken:
        suffix = f" ({number})"
        title, number = base[:31 - len(suffix)] + suffix, number + 1
    taken.add(title.lower())
    return title


def inspection_xlsx(session):
    """The audit as an editable workbook laid out like its PDFs: Overall first (details, score,
    charts, grading, every asset with its register details, findings, signatures), then one
    sheet per location in full (each asset's checks and photos, and the findings there)."""
    from datetime import datetime
    from openpyxl import Workbook
    from backend.xlsx_report import new_sheet, save
    settings = app_settings()
    summary = session.get("scoring") or summarize_score(session.get("items", []), settings)
    parts = [location_part(session, location, settings) for location in session_locations(session)]
    findings = session.get("findings") or []
    brand, media, printed = branding_settings(), MediaStore(config.DB_PATH), datetime.now().strftime("%Y-%m-%d %H:%M")
    workbook = Workbook()
    sheet = new_sheet(workbook, "Overall", brand, media, settings, first=True)
    sheet.letterhead("Facilities Audit Report", printed)
    sheet.audit_details(session)
    sheet.scorecard(summary, findings)
    sheet.charts(summary, parts, findings)
    sheet.location_grading(parts)
    sheet.asset_register(parts, asset_records(session.get("items", [])))
    sheet.findings(findings)
    sheet.signatures(session)
    taken = {"overall"}
    for part in parts:
        sheet = new_sheet(workbook, sheet_title(part["location"], taken), brand, media, settings)
        sheet.letterhead("Audit Report — Location", printed)
        sheet.banner(f"Location: {part['location']}")
        sheet.audit_details(session, part["location"])
        sheet.scorecard(part["summary"], part["findings"])
        sheet.checklist(part["items"], part["assets"])
        sheet.findings(part["findings"])
    return save(workbook)


def inspection_locations_pdf(session, locations):
    """The chosen locations of one audit in full: every check and remark, photos, and findings."""
    from backend.pdf_report import build_locations_report
    known = session_locations(session)
    chosen = list(dict.fromkeys(locations))
    if not chosen or any(location not in known for location in chosen):
        raise ValueError("Choose locations from this audit")
    settings = app_settings()
    return build_locations_report(session, [location_part(session, location, settings) for location in chosen],
                                  branding_settings(), MediaStore(config.DB_PATH), settings)


def location_audits(outlet, location, date_from="", date_to="", outlets=None):
    """Completed audits of an outlet that covered the location, each limited to it."""
    if outlets is not None and outlet not in outlets:
        return []
    clauses, params = ["outlet = ?", "status = 'Completed'"], [outlet]
    if date_from:
        clauses.append("audit_date >= ?")
        params.append(date_from)
    if date_to:
        clauses.append("audit_date <= ?")
        params.append(date_to)
    with connect() as db:
        ids = [row[0] for row in db.execute(f"SELECT id FROM inspection_sessions WHERE {' AND '.join(clauses)} ORDER BY audit_date, id", params)]
    settings = app_settings()
    audits = []
    for session_id in ids:
        session = inspection_session(session_id)
        items = [item for item in session.get("items", []) if (item.get("location") or session.get("zone")) == location]
        if not items:
            continue
        findings = [finding for finding in session.get("findings") or [] if finding.get("location") == location]
        audits.append({"session": session, "items": items, "findings": findings, "summary": summarize_score(items, settings),
                       "assets": asset_records(items)})
    return audits


def period_label(date_from, date_to):
    if date_from and date_to:
        return f"{date_from} to {date_to}"
    if date_from:
        return f"From {date_from}"
    if date_to:
        return f"Up to {date_to}"
    return "All dates"


def location_pdf(outlet, location, date_from="", date_to="", outlets=None):
    from backend.pdf_report import build_location_report
    audits = location_audits(outlet, location, date_from, date_to, outlets)
    return build_location_report(outlet, location, period_label(date_from, date_to), audits, branding_settings(),
                                 MediaStore(config.DB_PATH), app_settings())


def location_xlsx(outlet, location, date_from="", date_to="", outlets=None):
    """Every audit of one location, laid out like the location PDF: an overview of the audits,
    then a sheet per audit with its checks, photos, findings, and signatures there."""
    from datetime import datetime
    from openpyxl import Workbook
    from backend.xlsx_report import new_sheet, save
    audits = location_audits(outlet, location, date_from, date_to, outlets)
    settings = app_settings()
    brand, media, printed = branding_settings(), MediaStore(config.DB_PATH), datetime.now().strftime("%Y-%m-%d %H:%M")
    workbook = Workbook()
    sheet = new_sheet(workbook, "Overview", brand, media, settings, first=True)
    sheet.letterhead("Location Audit Report", printed)
    sheet.put(sheet.row, 1, f"{outlet} · {location}", size=11, bold=True, span=8, wrap=False)
    sheet.put(sheet.row + 1, 1, f"Period: {period_label(date_from, date_to)}", size=8, colour="5F6C66", span=8, wrap=False)
    sheet.row += 2
    sheet.heading("Audits of this location")
    rows, tones = [], {}
    for index, audit in enumerate(audits):
        session, summary = audit["session"], audit["summary"]
        rows.append([session.get("audit_date"), session.get("audit_ref") or "", session.get("auditor") or "", summary["total"],
                     summary["failed"], summary["score"], summary["rating"], len(audit["findings"])])
        tones[(index, 6)] = sheet.rating_tone(summary["score"])
    if rows:
        sheet.table(["Date", "Audit", "Auditor", "Checks", "Failed", "Score", "Rating", "Findings"], rows, [1] * 8, tones=tones)
    else:
        sheet.put(sheet.row, 1, "No completed audit covered this location in this period.", size=8, colour="5F6C66", span=8)
    taken = {"overview"}
    for audit in audits:
        session = audit["session"]
        sheet = new_sheet(workbook, sheet_title(session.get("audit_ref") or f"Audit {session.get('id')}", taken), brand, media, settings)
        sheet.letterhead("Location Audit Report", printed)
        sheet.banner(f"{session.get('audit_ref') or 'Audit'} · {session.get('audit_date') or ''} · {location}")
        sheet.audit_details(session, location)
        sheet.scorecard(audit["summary"], audit["findings"])
        sheet.checklist(audit["items"], audit["assets"])
        sheet.findings(audit["findings"])
        sheet.signatures(session)
    return save(workbook)
